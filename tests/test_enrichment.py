"""Fixtures fictícias para o contrato de enriquecimento; sem consulta comercial."""
import hashlib
import io
import json
import unittest
from openpyxl import load_workbook
from rota.enrichment import importar, CAMPOS
from rota.export import exportar
from rota.core import normalizar
from test_mvp import numero, source


def encode(data):return json.dumps(data,ensure_ascii=False,indent=2).encode()


class EnrichmentTests(unittest.TestCase):
    def test_escavador_multiple_sources_preserves_conflicts_and_raw_bytes(self):
        data={'numero_cnj':numero(),'extra_desconhecido':['preservar'], 'fontes':[
            {'processo_fonte_id':10,'grau':1,'status_predito':'INATIVO',
             'capa':{'valor_causa':{'valor':'100.00','moeda':'R$'}},
             'envolvidos':[{'nome':'Empresa TESTE','tipo_pessoa':'JURIDICA','polo':'ATIVO','cnpj':'00000000000000',
                            'advogados':[{'nome':'Advogado TESTE','oabs':[{'uf':'SC','numero':123}]}]},
                           {'nome':'Ente TESTE','polo':'PASSIVO'}]},
            {'processo_fonte_id':20,'grau':2,'capa':{'valor_causa':{'valor':'200.00'}},'envolvidos':[]}]}
        raw=b'\xef\xbb\xbf'+encode(data)
        c=importar(raw,'Escavador',{numero()},'admin@example.com')
        self.assertEqual(c['arquivo_original_utf8'].encode(),raw)
        self.assertEqual(c['sha256_arquivo'],hashlib.sha256(raw).hexdigest())
        self.assertEqual(c['original']['extra_desconhecido'],['preservar'])
        amounts=[o for o in c['observacoes'] if o['campo']=='valor_causa']
        self.assertEqual([o['grau'] for o in amounts],[1,2])
        for field in ('evidencia_transito','resultado_decisao'):
            self.assertEqual(next(o['status'] for o in c['observacoes'] if o['campo']==field),'nao_fornecido')
        self.assertEqual({o['campo'] for o in c['observacoes']},set(CAMPOS))

    def test_judit_inferred_company_documents_and_lawyers_are_not_confirmed(self):
        data={'code':numero(),'instance':1,'amount':123.45,'status':'FINALIZADO',
              'parties':[{'name':'Empresa TESTE','entity_type':'company','side':'Passive','was_inferred':True,
                          'documents':[{'document':'00000000000000','document_type':'cnpj'}],
                          'lawyers':[{'name':'Advogado TESTE','documents':[{'document':'SC123','document_type':'oab'}]}]}],
              'sentence':{'sentence_type':'Sentença','step_id':'step1'},
              'steps':[{'content':'Trânsito em julgado','step_id':'step2','lawsuit_instance':1}]}
        for body in (data,{'response_data':data}):
            c=importar(encode(body),'JUDIT',{numero()},'admin')
            for o in c['observacoes']:
                if o['campo'] in ('empresa','cnpj','advogado','polo_passivo'):self.assertEqual(o['status'],'inferido_fornecedor')
            self.assertEqual(next(o for o in c['observacoes'] if o['campo']=='resultado_decisao')['status'],'pendente_revisao')
            self.assertEqual(c['original'],body)

    def test_judit_status_is_not_transit_and_person_document_not_cnpj(self):
        data={'code':numero(),'status':'FINALIZADO','parties':[{'name':'Pessoa TESTE','entity_type':'person','main_document':'12345678901'}]}
        c=importar(encode(data),'JUDIT',{numero()},'admin')
        for f in ('evidencia_transito','cnpj','resultado_decisao'):
            self.assertEqual(next(o['status'] for o in c['observacoes'] if o['campo']==f),'nao_fornecido')

    def test_invalid_mismatched_and_credentials_rejected(self):
        for raw,provider in [(b'<html>','JUDIT'),(b'[]','JUDIT'),
            (encode({'code':numero(2)}),'JUDIT'),(encode({'code':'0'*20}),'JUDIT'),
            (encode({'code':numero(),'parties':{}}),'JUDIT'),
            (encode({'code':numero(),'unknown':{'Authorization':'SECRET'}}),'JUDIT'),
            (b' '* (5*1024*1024+1),'Escavador'),(encode({'code':numero()}),'Outro')]:
            with self.assertRaises(ValueError):importar(raw,provider,{numero()},'admin')

    def test_export_keeps_observations_separate_and_selected_only(self):
        data={'code':numero(),'amount':42,'parties':[], 'sentence':{'sentence_type':'=SUM(1,2)'+'a'*34000}}
        c=importar(encode(data),'JUDIT',{numero()},'admin')
        r={'id':'teste','criada':'2026-01-01','status':'Concluída','recebidos':2,'cnjs_unicos':2,
           'registros':[normalizar(source(),[]),normalizar(source(2),[])],'paginas':[], 'query':{},'duplicados':0,'complementos':[c]}
        original=json.dumps(r['registros'],sort_keys=True)
        w=load_workbook(io.BytesIO(exportar(r,[numero()])))
        self.assertIsNone(w['PROCESSOS']['R2'].value)
        self.assertIn('COMPLEMENTOS',w.sheetnames)
        self.assertGreater(w['FONTES_COMPLEMENTARES'].max_row,2)
        self.assertEqual(original,json.dumps(r['registros'],sort_keys=True))
        for sheet in w:
            for row in sheet:
                for cell in row:self.assertNotEqual(cell.data_type,'f')
        w=load_workbook(io.BytesIO(exportar(r,[numero(2)])))
        self.assertNotIn('COMPLEMENTOS',w.sheetnames)
