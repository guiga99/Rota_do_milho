import io,json,tempfile,unittest
from datetime import date
from pathlib import Path
import httpx
from openpyxl import load_workbook
from rota.core import DataJud,ConsultaError,Store,cnj_valido,normalizar,coletar
from rota.export import exportar

def numero(seq=1):
    base=f'{seq:07d}'+'2026'+'4'+'04'+'7000'
    return base[:7]+f'{98-int(base+"00")%97:02d}'+base[7:]
def source(seq=1,grau='G1'):
    return {'id':f'r{seq}-{grau}','numeroProcesso':numero(seq),'tribunal':'TRF4','grau':grau,'dataAjuizamento':'20260101120000','classe':{'nome':'=formula'},'orgaoJulgador':{'nome':'Órgão'},'assuntos':[{'codigo':6060,'nome':'Folha'}],'movimentos':[{'codigo':848,'nome':'Trânsito em Julgado','dataHora':'2026-02-01T12:00:00Z'}]}
def response(sources):return {'hits':{'total':{'value':len(sources),'relation':'eq'},'hits':[{'_source':s,'sort':[s['dataAjuizamento'],s['id']]} for s in sources]}}

class MVPTests(unittest.TestCase):
    def test_cnj_and_missing(self):
        self.assertTrue(cnj_valido(numero()));self.assertFalse(cnj_valido('0'*20));self.assertFalse(cnj_valido(123))
        s=source();s['dataAjuizamento']='invalid';s['assuntos']=[]
        r=normalizar(s,[6060]);self.assertEqual(r['qualidade'],'Com ressalvas');self.assertIn('Assunto incompatível',r['ressalvas'])
        self.assertTrue(r['indicio_transito'])
    def test_periodo(self):
        self.assertIn('fora do período',normalizar(source(),[6060],date(2026,2,1),None)['ressalvas'])
    def test_auth_sanitized(self):
        api=DataJud('secret',httpx.MockTransport(lambda req:httpx.Response(401,text='secret')))
        with self.assertRaises(ConsultaError) as c:api.teste()
        self.assertNotIn('secret',str(c.exception));api.close()
    def test_invalid_200(self):
        for payload in ['<html>',json.dumps({'hits':{}}),json.dumps({'hits':{'hits':[{'_source':{'tribunal':'TRF3'}}]}})]:
            api=DataJud('key',httpx.MockTransport(lambda req:httpx.Response(200,text=payload)))
            with self.assertRaises(ConsultaError):api.teste()
            api.close()
    def test_empty(self):
        api=DataJud('key',httpx.MockTransport(lambda req:httpx.Response(200,json=response([]))))
        with tempfile.TemporaryDirectory() as d:
            r=coletar(api,Store(Path(d)/'db'),[6060],10)
            self.assertEqual(r['status'],'Concluída');self.assertEqual(r['cnjs_unicos'],0)
            self.assertEqual(load_workbook(io.BytesIO(exportar(r)))['PROCESSOS'].max_row,1)
        api.close()
    def test_partial_and_history(self):
        count=0
        def request(req):
            nonlocal count
            count+=1
            if count==1:return httpx.Response(200,json=response([source(),source(1,'G2')]))
            return httpx.Response(401)
        api=DataJud('key',httpx.MockTransport(request))
        with tempfile.TemporaryDirectory() as d:
            store=Store(Path(d)/'db');r=coletar(api,store,[6060],2)
            self.assertEqual(r['status'],'Parcial');self.assertEqual(r['cnjs_unicos'],1)
            self.assertEqual(len(r['registros']),2);self.assertEqual(store.obter(r['id']),r)
        api.close()
    def test_export_and_repeat(self):
        api=DataJud('key',httpx.MockTransport(lambda req:httpx.Response(200,json=response([source(),source(),source(2)]))))
        with tempfile.TemporaryDirectory() as d:
            store=Store(Path(d)/'db');r=coletar(api,store,[6060],5)
            self.assertEqual(r['duplicados'],1);self.assertEqual(r['cnjs_unicos'],2)
            binary=exportar(r,[numero()]);w=load_workbook(io.BytesIO(binary))
            self.assertEqual(w['PROCESSOS'].max_row,2);self.assertEqual(w['PROCESSOS']['A2'].data_type,'s')
            self.assertTrue(w['PROCESSOS']['E2'].value.startswith("'="));self.assertEqual(w['PROCESSOS']['E2'].data_type,'s')
            self.assertIsNone(w['PROCESSOS']['M2'].value)
            r2=coletar(api,store,[6060],5);self.assertEqual(len(store.listar()),2);self.assertNotEqual(r['id'],r2['id'])
        api.close()

if __name__=='__main__':unittest.main()
