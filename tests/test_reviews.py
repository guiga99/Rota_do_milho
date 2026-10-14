import copy
import io
import tempfile
import unittest
import os
from unittest.mock import patch
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from openpyxl import load_workbook
from rota.core import Store, normalizar
from rota.reviews import alvos, registrar, ultimas, linhas_revisoes
from rota.export import exportar
from test_mvp import source, numero


def fixture():
    return {'id':'fixture-reviews','criada':'2026-01-01','status':'Concluída','recebidos':2,'cnjs_unicos':2,
            'registros':[normalizar(source(),[]),normalizar(source(2),[])],'paginas':[],'query':{},'duplicados':0}


def args(result,email='colega@example.com',estado='Confirmado'):
    t=next(t for t in alvos(result) if t['campo']=='classe.nome')
    return dict(alvo_id=t['alvo_id'],estado=estado,comentario='Conferido na fonte de teste',proposta='',evidencia='Documento TESTE',
                identity={'email':email,'role':'operator','mode':'google'})


class ReviewTests(unittest.TestCase):
    def test_google_operator_saves_with_authenticated_identity(self):
        from streamlit.testing.v1 import AppTest
        with tempfile.TemporaryDirectory() as d,patch.dict(os.environ,{'ROTA_DATA_DIR':d}),patch('rota.access.proteger',return_value={'email':'colega@example.com','role':'operator','mode':'google'}):
            store=Store(Path(d)/'consultas.sqlite3');r=fixture();store.salvar(r)
            app=AppTest.from_file(str(Path(__file__).parents[1]/'app.py'))
            app.session_state['result']=r;app.run()
            self.assertEqual(len(app.exception),0)
            self.assertIsNone(app.radio[0].value)
            app.radio[0].set_value('Confirmado')
            app.text_area[0].set_value('Conferência TESTE')
            app.text_area[2].set_value('Documento TESTE')
            next(b for b in app.button if b.label=='Salvar revisão do campo').click().run()
            self.assertEqual(len(app.exception),0)
            saved=store.obter(r['id'])
            self.assertEqual(saved['revisoes'][0]['revisor'],'colega@example.com')
            self.assertEqual(saved['registros'],r['registros'])

    def test_source_preserved_and_identity_from_server(self):
        r=fixture();original=copy.deepcopy(r);updated=registrar(r,**args(r))
        self.assertEqual(r,original);self.assertEqual(updated['registros'],r['registros'])
        v=updated['revisoes'][0]
        self.assertEqual(v['revisor'],'colega@example.com');self.assertEqual(v['valor_original'],'=formula')
        self.assertTrue(v['revisada_utc']);self.assertEqual(v['estado'],'Confirmado')

    def test_absent_cannot_be_confirmed_and_correction_requires_evidence(self):
        r=fixture();kw=args(r)
        kw['alvo_id']=next(t['alvo_id'] for t in alvos(r) if t['campo']=='complementar.cnpj')
        with self.assertRaises(ValueError):registrar(r,**kw)
        kw['estado']='Informação insuficiente';kw['evidencia']=''
        self.assertEqual(registrar(r,**kw)['revisoes'][0]['estado'],'Informação insuficiente')
        for change in [{'estado':None},{'comentario':''},{'estado':'Corrigir'},{'estado':'Corrigir','proposta':'TESTE','evidencia':''},
                       {'identity':{'email':'fake','role':'guest','mode':'google'}},{'comentario':'a\x00b'}]:
            with self.assertRaises(ValueError):registrar(r,**{**args(r),**change})

    def test_independent_reviewers_and_version_history(self):
        r=fixture();r=registrar(r,**args(r));first=r['revisoes'][0]['id']
        r=registrar(r,**args(r,'outro@example.com','Informação insuficiente'))
        r=registrar(r,**{**args(r),'estado':'Corrigir','proposta':'Classe TESTE','anterior':first})
        self.assertEqual(len(r['revisoes']),3);self.assertEqual(len(ultimas(r)),2)
        self.assertEqual(r['revisoes'][-1]['substitui'],first)
        with self.assertRaises(ValueError):registrar(r,**args(r),anterior=first)

    def test_changed_snapshot_invalidates_old_target_but_preserves_history(self):
        r=fixture();kw=args(r);r=registrar(r,**kw)
        r['registros'][0]['original']['classe']['nome']='Outro valor'
        with self.assertRaises(ValueError):registrar(r,**kw,anterior=r['revisoes'][0]['id'])
        self.assertFalse(list(linhas_revisoes(r))[0]['Origem vigente'])

    def test_transactional_concurrent_reviews_and_reload(self):
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'db';store=Store(path);r=fixture();store.salvar(r)
            def save(i):return store.revisar(r['id'],**args(r,f'colega{i}@example.com'))
            with ThreadPoolExecutor(max_workers=4) as pool:list(pool.map(save,range(4)))
            saved=Store(path).obter(r['id'])
            self.assertEqual(len(saved['revisoes']),4);self.assertEqual(saved['registros'],r['registros'])
            store.salvar(r)
            self.assertEqual(len(store.obter(r['id'])['revisoes']),4)
            with self.assertRaises(ValueError):store.revisar(r['id'],**args(r,'colega0@example.com'))
            self.assertEqual(len(store.obter(r['id'])['revisoes']),4)

    def test_export_selection_audit_and_formula_safety(self):
        r=fixture();r=registrar(r,**{**args(r),'comentario':'=1+1'})
        w=load_workbook(io.BytesIO(exportar(r,[numero()])))
        sheet=w['REVISOES_CAMPOS'];headers=[c.value for c in sheet[1]]
        row=dict(zip(headers,[c.value for c in sheet[2]]))
        self.assertEqual(row['CNJ'],numero());self.assertEqual(row['Revisor'],'colega@example.com')
        self.assertTrue(row['Comentário'].startswith("'="))
        self.assertNotIn('REVISOES_CAMPOS',load_workbook(io.BytesIO(exportar(r,[numero(2)]))).sheetnames)
        for s in w:
            for rows in s:
                for c in rows:self.assertNotEqual(c.data_type,'f')
