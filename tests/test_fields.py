import io,tempfile,unittest
from datetime import date
from pathlib import Path
from openpyxl import load_workbook
from rota.core import DataJud,Store,coletar
from rota.fields import inventario,periodo_cinco_anos,valores
from rota.export import exportar
from test_mvp import source,response
import httpx

class AllFieldsTests(unittest.TestCase):
    def test_five_years_and_leap_day(self):
        self.assertEqual(periodo_cinco_anos(date(2024,2,29)),(date(2019,2,28),date(2024,2,29)))
        self.assertEqual(periodo_cinco_anos(date(2026,10,13))[0],date(2021,10,13))
    def test_inventory_is_per_record_not_per_array_element(self):
        a={'original':{'extra':[{'v':None},{'v':12}], 'empty':[]}}
        b={'original':{'extra':[{'v':'012'}]}}
        fields={x['campo']:x for x in inventario([a,b])}
        self.assertEqual(fields['extra[].v']['registros_presentes'],2)
        self.assertEqual(fields['extra[].v']['ocorrencias'],3)
        self.assertEqual(fields['extra[].v']['nulos'],1)
        self.assertEqual(fields['empty']['registros_presentes'],1)
        self.assertEqual(list(valores({'empty':[], 'null':None})),[('empty','array',[]),('null','null',None)])
    def test_export_preserves_unknown_fields_and_long_text(self):
        s=source();s['novo']={'nome':'=externo','texto':'á'*33000,'codigo':'0012','nulo':None}
        api=DataJud('key',httpx.MockTransport(lambda req:httpx.Response(200,json=response([s]))))
        with tempfile.TemporaryDirectory() as d:
            r=coletar(api,Store(Path(d)/'db'),[6060],2)
            w=load_workbook(io.BytesIO(exportar(r)))
            rows=list(w['DADOS_ORIGINAIS'].values)[1:]
            text=[row for row in rows if row[2]=='novo.texto']
            self.assertEqual(''.join(row[6] for row in text),'á'*33000)
            code=next(row for row in rows if row[2]=='novo.codigo')
            self.assertEqual(code[6],'0012')
            formula=next(row for row in rows if row[2]=='novo.nome')
            self.assertEqual(formula[6],"'=externo")
            self.assertEqual(next(row for row in rows if row[2]=='novo.nulo')[3],'null')
            self.assertIn('novo.texto',[row[0] for row in w['CAMPOS_DISPONIVEIS'].values][1:])
        api.close()

class LeanTableTests(unittest.TestCase):
    def test_configurable_table_preserves_full_original(self):
        from rota.fields import extrair,tabela_principal
        s=source();s['assuntos'].append({'codigo':123,'nome':'Outro'})
        self.assertEqual(extrair(s,'assuntos[].nome'),['Folha','Outro'])
        self.assertIsNone(extrair(s,'unknown.field'))
        api=DataJud('key',httpx.MockTransport(lambda req:httpx.Response(200,json=response([s]))))
        with tempfile.TemporaryDirectory() as d:
            r=coletar(api,Store(Path(d)/'db'),[6060],2)
            w=load_workbook(io.BytesIO(exportar(r,campos=['numeroProcesso','classe.nome'])))
            self.assertEqual(list(w['TABELA_PRINCIPAL'].values)[0],('Número CNJ','Classe','Qualidade técnica'))
            self.assertIn('DADOS_ORIGINAIS',w.sheetnames)
            self.assertIn('assuntos[1].nome',[row[2] for row in list(w['DADOS_ORIGINAIS'].values)[1:]])
            self.assertEqual(r['registros'][0]['original'],s)
        api.close()
