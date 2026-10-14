"""Fluxo real no navegador com arquivos sintéticos e banco local descartável."""
import hashlib
import json
import os
import re
import shutil
import socket
import subprocess
import sys
import tempfile
import time
import urllib.request
from pathlib import Path
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

ROOT=Path(__file__).parents[1]
sys.path.insert(0,str(ROOT))
from rota.core import Store, normalizar
from test_mvp import numero, source

with tempfile.TemporaryDirectory(prefix='rota-browser-') as d:
    root=Path(d);store=Store(root/'consultas.sqlite3')
    result={'id':'browser-fixture','criada':'2026-01-01','status':'Concluída','recebidos':1,'cnjs_unicos':1,
            'registros':[normalizar(source(),[])],'paginas':[],'query':{},'duplicados':0}
    store.salvar(result)
    payload={'code':numero(),'instance':1,'amount':123.45,'parties':[
        {'name':'EMPRESA TESTE','entity_type':'company','side':'Active','was_inferred':True,
         'documents':[{'document':'00000000000000','document_type':'cnpj'}]}]}
    raw=json.dumps(payload,ensure_ascii=False,indent=2).encode()
    upload=root/'fixture.json';upload.write_bytes(raw)
    wrong=root/'cnj-diferente.json';wrong.write_text(json.dumps({**payload,'code':numero(2)}))
    with socket.socket() as s:s.bind(('127.0.0.1',0));port=s.getsockname()[1]
    env={**os.environ,'ROTA_DATA_DIR':d,'ROTA_AUTH_MODE':'local','DATAJUD_API_KEY':''}
    with (root/'streamlit.log').open('w') as log:
        process=subprocess.Popen([sys.executable,'-m','streamlit','run','app.py',
            '--server.address','127.0.0.1','--server.port',str(port),'--server.headless','true'],cwd=ROOT,env=env,stdout=log,stderr=log)
        try:
            url=f'http://127.0.0.1:{port}'
            for _ in range(40):
                try:
                    with urllib.request.urlopen(url+'/_stcore/health',timeout=1) as r:
                        if r.status==200:break
                except OSError:time.sleep(.25)
            else:raise RuntimeError('Aplicação de teste não iniciou.')
            with sync_playwright() as p:
                browser=p.chromium.launch(executable_path=shutil.which('chromium'),headless=True,args=['--no-sandbox'])
                page=browser.new_page();page.goto(url)
                page.get_by_role('tab',name='Histórico',exact=True).click()
                page.get_by_role('button',name='Abrir consulta',exact=True).click()
                page.get_by_role('tab',name='Resultados',exact=True).click()
                title=page.get_by_text('Complementar informações com JSON de fornecedor',exact=True)
                title.click()
                combo=page.get_by_test_id('stSelectbox').filter(has_text='Formato da resposta').get_by_role('combobox')
                combo.click();page.get_by_role('option',name='JUDIT',exact=True).click()
                page.locator('input[type=file]').set_input_files(str(wrong))
                page.get_by_role('button',name='Importar JSON e registrar fonte',exact=True).click()
                page.get_by_text('O CNJ é inválido ou não pertence à consulta aberta.',exact=True).wait_for()
                assert not store.obter(result['id']).get('complementos')
                page.locator('input[type=file]').set_input_files(str(upload))
                page.get_by_role('button',name='Importar JSON e registrar fonte',exact=True).click()
                page.get_by_role('button',name='Baixar XLSX verificado',exact=True).wait_for()
                for _ in range(30):
                    saved=store.obter(result['id'])
                    if saved.get('complementos'):break
                    time.sleep(.1)
                else:raise AssertionError('Importação não foi persistida.')
                page.get_by_text(re.compile('Observações não foram aprovadas juridicamente')).wait_for(state='attached')
                page.wait_for_function('document.querySelector(\'[data-testid="stApp"]\')?.getAttribute("data-test-script-state") === "notRunning"')
                assert saved['registros']==result['registros']
                page.get_by_text('Validar campos com a equipe',exact=True).click()
                radio=page.get_by_role('radio',name='Confirmado',exact=True)
                radio.locator('..').click(timeout=5000)
                assert page.get_by_role('radio',name='Confirmado',exact=True).is_checked()
                page.get_by_label('Comentário da revisão',exact=True).fill('Campos conferidos na fixture TESTE.')
                page.get_by_label('Fonte ou evidência da conferência',exact=True).fill('Documento sintético TESTE')
                page.get_by_role('button',name='Salvar revisão do campo',exact=True).click()
                page.get_by_text('Revisão registrada com revisor e data. O valor de origem foi preservado.',exact=True).wait_for(state='attached')
                page.wait_for_function('document.querySelector(\'[data-testid="stApp"]\')?.getAttribute("data-test-script-state") === "notRunning"')
                reviewed=store.obter(result['id'])
                assert reviewed['revisoes'][0]['revisor']=='operador local'
                assert reviewed['registros']==result['registros']
                with page.expect_download() as dl:page.get_by_role('button',name='Baixar XLSX verificado',exact=True).click()
                path=root/'download.xlsx';dl.value.save_as(path)
                workbook=load_workbook(path)
                assert 'COMPLEMENTOS' in workbook.sheetnames
                assert workbook['REVISOES_CAMPOS'].max_row==2
                assert workbook['PROCESSOS']['A2'].value==numero()
                assert workbook['PROCESSOS']['M2'].value is None
                with page.expect_download() as dl:page.get_by_role('button',name='Baixar evidências JSON',exact=True).click()
                evidence=root/'evidence.json';dl.value.save_as(evidence)
                c=json.loads(evidence.read_bytes())['complementos'][0]
                assert c['arquivo_original_utf8'].encode()==raw
                assert c['sha256_arquivo']==hashlib.sha256(raw).hexdigest()
                assert any(o['status']=='inferido_fornecedor' for o in c['observacoes'])
                browser.close()
            print('OK: navegador rejeita CNJ diferente, importa JSON JUDIT, salva revisão atribuída e baixa XLSX/JSON com histórico/original/hash preservados. Dados sintéticos.')
        finally:
            process.terminate()
            try:process.wait(timeout=10)
            except subprocess.TimeoutExpired:process.kill();process.wait()
