import io,json
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill
from .core import ENDPOINT
from .fields import inventario, valores, tabela_principal
from .enrichment import linhas
from .reviews import linhas_revisoes

def texto(v):
    if isinstance(v,str) and v.startswith(('=','+','-','@')):return "'"+v
    return v

def exportar(result,selecionados=None,campos=None):
    records=[r for r in result['registros'] if r['qualidade']!='Não exportável' and (selecionados is None or r['cnj'] in selecionados)]
    w=Workbook();w.remove(w.active)
    def aba(nome,headers,rows):
        s=w.create_sheet(nome);s.append(headers)
        for row in rows:
            if any(isinstance(v,str) and len(v)>32767 for v in row):raise ValueError('Conteúdo excede limite de célula do Excel; consulte a evidência JSON.')
            s.append([texto(v) for v in row])
        s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
        for c in s[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='214E34')
        for col in s.columns:s.column_dimensions[col[0].column_letter].width=min(60,max(18,max(len(str(c.value or '')) for c in col)+2))
        if s.max_row>1048576:raise ValueError('Exportação excede limite de linhas do Excel.')
        return s
    aba('RESUMO',['Item','Valor'],[['Consulta',result['id']],['Coleta UTC',result['criada']],['Status',result['status']],['Registros recebidos',result['recebidos']],['CNJs únicos na amostra',result['cnjs_unicos']],['CNJs exportados',len({r['cnj'] for r in records})],['Registros de origem exportados',len(records)],['Seleção local',json.dumps(selecionados,ensure_ascii=False) if selecionados is not None else 'Todos os tecnicamente exportáveis'],['Revisão jurídica','Pendente; indício textual não confirma decisão favorável'],['Erro',result.get('erro')]])
    if campos:
        lean=tabela_principal(records,campos)
        headers=list(lean[0]) if lean else [*campos,'Qualidade técnica']
        aba('TABELA_PRINCIPAL',headers,[[row.get(k) for k in headers] for row in lean])
    p=aba('PROCESSOS',['Número CNJ','ID origem','Tribunal','Grau','Classe','Órgão','Assuntos','Ajuizamento original','Qualidade técnica','Ressalvas','Indício textual de trânsito','Evidência trânsito','Empresa','CNPJ','Advogado','OAB','Polo passivo','Valor da causa','Resultado decisão'],[[r['cnj'],r['origem_id'],r['tribunal'],r['grau'],r['classe'],r['orgao'],r['assuntos'],r['ajuizamento'],r['qualidade'],r['ressalvas'],r['indicio_transito'],r['evidencia_transito'],None,None,None,None,None,None,None] for r in records])
    for c in list(p.columns)[0][1:]:c.number_format='@'
    aba('MOVIMENTACOES',['Número CNJ','ID origem','Código','Nome','Data original'],[[r['cnj'],r['origem_id'],m.get('codigo'),m.get('nome'),m.get('dataHora')] for r in records for m in r['movimentos']])
    aba('CAMPOS_DISPONIVEIS', ['Campo original', 'Tipos', 'Registros presentes', 'Total registros', 'Ocorrências', 'Nulos'], [[x['campo'],x['tipos'],x['registros_presentes'],x['total_registros'],x['ocorrencias'],x['nulos']] for x in inventario(records)])
    def originais():
        for r in records:
            for path,kind,value in valores(r['original']):
                cell=json.dumps(value,ensure_ascii=False) if kind in ('array','object') else value
                if isinstance(cell,str):
                    # JSON completo preserva a fonte; fragmentação evita truncar células longas.
                    chunks=[cell[i:i+32000] for i in range(0,len(cell),32000)] or ['']
                else: chunks=[cell]
                for index,chunk in enumerate(chunks,1): yield [r['cnj'],r['origem_id'],path,kind,index,len(chunks),chunk]
    aba('DADOS_ORIGINAIS', ['Número CNJ', 'ID origem', 'Caminho com índice', 'Tipo', 'Fragmento', 'Total fragmentos', 'Valor original seguro'], originais())
    aba('FONTES',['Página','Fonte','Coleta UTC','SHA256 JSON canônico'],[[i,result.get('endpoint',ENDPOINT),p['coletada'],p['sha256']] for i,p in enumerate(result['paginas'],1)])
    complementos=[c for c in result.get('complementos',[]) if c['cnj'] in {r['cnj'] for r in records}]
    if complementos:
        def observacoes():
            for row in linhas(complementos):
                cell=row['Valor informado']
                chunks=[cell[i:i+32000] for i in range(0,len(cell),32000)] if cell else [cell]
                for i,chunk in enumerate(chunks,1):yield [*list(row.values())[:3],chunk,*list(row.values())[4:],i,len(chunks)]
        aba('COMPLEMENTOS',['CNJ','Fornecedor','Campo','Valor informado','Estado','Grau','ID fonte','Caminho JSON','Importação UTC','SHA256 arquivo','Fragmento','Total fragmentos'],observacoes())
        rows=[]
        for c in complementos:
            raw=json.dumps(c['original'],ensure_ascii=False)
            chunks=[raw[i:i+32000] for i in range(0,len(raw),32000)]
            for i,chunk in enumerate(chunks,1):rows.append([c['cnj'],c['fornecedor'],c['importada_utc'],c['sha256_arquivo'],c['documentacao'],c['autenticidade'],i,len(chunks),chunk])
        aba('FONTES_COMPLEMENTARES',['CNJ','Fornecedor','Importação UTC','SHA256 arquivo original','Documentação','Autenticidade','Fragmento','Total fragmentos','JSON reserializado'],rows)
    revisoes=[r for r in linhas_revisoes(result) if r['CNJ'] in {x['cnj'] for x in records}]
    if revisoes:
        headers=list(revisoes[0])
        def revisadas():
            for r in revisoes:
                value=r['Valor original'];chunks=[value[i:i+32000] for i in range(0,len(value),32000)] or ['']
                for i,chunk in enumerate(chunks,1):yield [chunk if k=='Valor original' else r[k] for k in headers]+[i,len(chunks)]
        aba('REVISOES_CAMPOS',headers+['Fragmento','Total fragmentos'],revisadas())
    aba('PENDENCIAS',['Item','Limitação'],[['Campos ausentes','Colunas de partes/valor/decisão em PROCESSOS preservam a fonte; complementos e revisões constam em abas próprias.'],['Revisões','Conferência por campo não é aceite jurídico do processo. Correções são propostas, sem sobrescrever a fonte.'],['Assuntos','Códigos iniciais observados no piloto; hierarquia e catálogo ainda precisam de revisão jurídica.'],['Trânsito','Indício pelo nome do movimento; regras por códigos e revisão jurídica pendentes.'],['Abrangência','Amostra limitada; registros de outros graus fora da amostra podem existir.']])
    aba('EXECUCAO',['Consulta','Query JSON','Status','Recebidos','Duplicados','Total API'],[[result['id'],json.dumps(result['query'],ensure_ascii=False),result['status'],result['recebidos'],result['duplicados'],json.dumps(result.get('total_api'),ensure_ascii=False)]])
    aba('VALIDACAO',['Número CNJ','ID origem','Qualidade','Motivo'],[[r['cnj'],r['origem_id'],r['qualidade'],r['ressalvas']] for r in result['registros'] if selecionados is None or r['cnj'] in selecionados])
    b=io.BytesIO();w.save(b);data=b.getvalue();v=load_workbook(io.BytesIO(data))
    if v['PROCESSOS'].max_row-1!=len(records):raise ValueError('Contagem divergente na exportação.')
    expected=[[None if x == "" else x for x in row] for row in p.values]
    actual=[list(row) for row in v['PROCESSOS'].values]
    if expected!=actual:raise ValueError('Conteúdo divergente ao reabrir Excel.')
    return data
